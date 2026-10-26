# -*- coding: utf-8 -*-
"""
Sinh kịch bản thuyết trình từ noi-dung.json (dùng chung với build_pptx.js):
  - KICH-BAN-THUYET-TRINH.pdf  (qua .docx rồi LibreOffice chuyển sang PDF)
  - KICH-BAN-THUYET-TRINH.md   (bản đọc nhanh trong repo)

    cd tools/thuyet-trinh && python3 build_kich_ban.py
"""
import json
import os
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DAY = os.path.dirname(os.path.abspath(__file__))
GOC = os.path.dirname(os.path.dirname(DAY))
ND = json.load(open(os.path.join(DAY, "noi-dung.json"), encoding="utf-8"))
PDF = os.path.join(GOC, "KICH-BAN-THUYET-TRINH.pdf")
MD = os.path.join(GOC, "KICH-BAN-THUYET-TRINH.md")

NAVY = RGBColor(0x13, 0x29, 0x4B)
CAM = RGBColor(0xE8, 0x62, 0x2A)
XAM = RGBColor(0x57, 0x60, 0x6A)
MAU_NGUOI = {"Tín": "E8622A", "Nhân": "1F6FEB", "Khôi": "8250DF", "Vy": "1A7F37"}


def phut(giay):
    m, s = divmod(int(giay), 60)
    return "%d:%02d" % (m, s)


def to_nen(o, hex_):
    tc = o._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_)
    tc.append(sh)


def bang(doc, tieu_de, dong, rong):
    t = doc.add_table(rows=1, cols=len(tieu_de))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(tieu_de):
        o = t.rows[0].cells[i]
        o.text = ""
        r = o.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        to_nen(o, "13294B")
    for k, d in enumerate(dong):
        hang = t.add_row().cells
        for i, v in enumerate(d):
            hang[i].text = ""
            hang[i].paragraphs[0].add_run(str(v))
            if k % 2:
                to_nen(hang[i], "EEF2F7")
    # LibreOffice đọc độ rộng ở <w:tblGrid>, Word đọc ở từng ô — đặt cả hai.
    t.autofit = False
    for i, w in enumerate(rong):
        t.columns[i].width = Cm(w)
    for hang in t.rows:
        for i, w in enumerate(rong):
            hang.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def tieu_de(doc, chu, cap):
    h = doc.add_heading(chu, level=cap)
    for r in h.runs:
        r.font.color.rgb = NAVY
    return h


def tao_docx(duong):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2)
        sec.left_margin = sec.right_margin = Cm(2.2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(11.5)
    for ten in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        doc.styles[ten].font.name = "Cambria"

    # chân trang: số trang
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Kịch bản thuyết trình · IE402 · Nhóm 1 · trang ")
    r.font.size = Pt(9)
    r.font.color.rgb = XAM
    for kieu, chu in (("begin", None), (None, "PAGE"), ("end", None)):
        rr = p.add_run()
        rr.font.size = Pt(9)
        if kieu:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kieu)
            rr._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = chu
            rr._r.append(it)

    # ---------------- bìa
    t = doc.add_paragraph()
    r = t.add_run("KỊCH BẢN THUYẾT TRÌNH")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = NAVY
    r.font.name = "Cambria"
    t = doc.add_paragraph()
    r = t.add_run(ND["tieu_de"])
    r.font.size = Pt(15)
    r.font.color.rgb = CAM
    doc.add_paragraph(ND["mon"] + " · GVHD: " + ND["gvhd"] + " · " + ND["nhom"])

    tong = sum(s["giay"] for s in ND["slides"])
    demo = next(s["giay"] for s in ND["slides"] if s["tieu_de"] == "Demo trực tiếp")
    p = doc.add_paragraph()
    p.add_run("Thời lượng: ").bold = True
    p.add_run("khoảng %s phút trình bày + %s phút demo = %s phút, chưa tính hỏi đáp. "
              "Tệp slide: THUYET-TRINH-IE402-Cham-cong-dinh-vi-3D.pptx (%d slide); lời nói "
              "dưới đây cũng nằm trong phần ghi chú (Notes) của từng slide."
              % (round((tong - demo) / 60), round(demo / 60), round(tong / 60), len(ND["slides"])))

    # ---------------- phân công lời
    tieu_de(doc, "1. Ai nói slide nào", 1)
    theo_nguoi = {}
    for i, s in enumerate(ND["slides"], start=1):
        theo_nguoi.setdefault(s["nguoi"], []).append((i, s))
    dong = []
    for tv in ND["thanh_vien"]:
        ten_goi = tv["ten"].split()[-1]
        ds = theo_nguoi.get(ten_goi, [])
        so = ", ".join(str(i) for i, _ in ds)
        dong.append([tv["ten"] + " (" + tv["mssv"] + ")", tv["vai"], so,
                     phut(sum(s["giay"] for _, s in ds))])
    bang(doc, ["Thành viên", "Vai trò", "Slide", "Thời gian"], dong, [4.6, 4.6, 4.4, 2.0])
    doc.add_paragraph("Chuyển người: người trước nói “Em xin mời bạn …”, người sau bước lên rồi mới "
                      "bắt đầu nói. Lời ở cuối mỗi đoạn dưới đây đã có sẵn câu chuyển.")

    # ---------------- lời từng slide
    tieu_de(doc, "2. Lời nói theo từng slide", 1)
    da = 0
    for i, s in enumerate(ND["slides"], start=1):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.keep_with_next = True
        r = h.add_run("Slide %d · %s" % (i, s["tieu_de"]))
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = NAVY
        r.font.name = "Cambria"
        m = doc.add_paragraph()
        m.paragraph_format.keep_with_next = True
        r = m.add_run(s["nguoi"])
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(MAU_NGUOI.get(s["nguoi"], "57606A"))
        r2 = m.add_run("   ·   ~%s   ·   mốc %s → %s" % (phut(s["giay"]), phut(da), phut(da + s["giay"])))
        r2.font.color.rgb = XAM
        r2.font.size = Pt(10)
        da += s["giay"]
        doc.add_paragraph(s["loi"])

    # ---------------- demo
    tieu_de(doc, "3. Kịch bản demo trực tiếp", 1)
    tieu_de(doc, "Chuẩn bị 15 phút trước", 2)
    for c in ND["demo"]["chuan_bi"]:
        doc.add_paragraph(c, style="List Bullet")
    tieu_de(doc, "Sáu bước", 2)
    for k, b in enumerate(ND["demo"]["buoc"], start=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        r = p.add_run("Bước %d · %s  (~%s)" % (k, b["ten"], phut(b["giay"])))
        r.bold = True
        r.font.color.rgb = CAM
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run("Làm: ").bold = True
        p.add_run(b["lam"])
        p = doc.add_paragraph()
        p.add_run("Nói: ").bold = True
        p.add_run(b["noi"])
    tieu_de(doc, "Khi có sự cố", 2)
    bang(doc, ["Tình huống", "Xử lý"], ND["demo"]["su_co"], [4.5, 11.1])

    # ---------------- hỏi đáp
    tieu_de(doc, "4. Câu hỏi dễ gặp và người trả lời", 1)
    bang(doc, ["Câu hỏi", "Người", "Ý chính"], ND["hoi_dap"], [4.6, 1.4, 9.6])

    tieu_de(doc, "5. Lưu ý khi trình bày", 1)
    for c in [
        "Không đọc slide — slide đã có chữ, người nói bổ sung “vì sao”.",
        "Dừng khoảng 2 giây sau con số quan trọng, nhất là “62,78 mét” và “13,5 tầng”.",
        "Không biết thì nói thật: “Phần đó nhóm chưa khảo sát kỹ, em xin ghi nhận ạ.”",
        "Thầy hỏi giữa chừng: trả lời ngắn, xin trả lời đầy đủ ở phần hỏi đáp cuối.",
        "Số trên màn hình demo có thể khác ảnh trên slide (khác phiên chạy) — logic giống nhau, "
        "không cần lúng túng.",
    ]:
        doc.add_paragraph(c, style="List Bullet")
    doc.save(duong)


def tao_md():
    d = ["# Kịch bản thuyết trình — Đồ án IE402", "",
         "> Sinh tự động từ `tools/thuyet-trinh/noi-dung.json` — sửa nội dung ở tệp đó rồi chạy "
         "`npm run build` trong `tools/thuyet-trinh/` để cập nhật cả slide (ghi chú), bản PDF và tệp này.",
         "", "**" + ND["tieu_de"] + "** · " + ND["mon"] + " · GVHD " + ND["gvhd"], ""]
    d += ["## Lời nói theo từng slide", ""]
    for i, s in enumerate(ND["slides"], start=1):
        d += ["### Slide %d · %s — %s (~%s)" % (i, s["tieu_de"], s["nguoi"], phut(s["giay"])), "", s["loi"], ""]
    d += ["## Demo trực tiếp", "", "**Chuẩn bị**", ""]
    d += ["- " + c for c in ND["demo"]["chuan_bi"]] + [""]
    for k, b in enumerate(ND["demo"]["buoc"], start=1):
        d += ["**%d · %s** (~%s)" % (k, b["ten"], phut(b["giay"])), "",
              "- Làm: " + b["lam"], "- Nói: " + b["noi"], ""]
    d += ["**Khi có sự cố**", "", "| Tình huống | Xử lý |", "|---|---|"]
    d += ["| %s | %s |" % tuple(r) for r in ND["demo"]["su_co"]] + [""]
    d += ["## Câu hỏi dễ gặp", "", "| Câu hỏi | Người | Ý chính |", "|---|---|---|"]
    d += ["| %s | %s | %s |" % tuple(r) for r in ND["hoi_dap"]] + [""]
    open(MD, "w", encoding="utf-8").write("\n".join(d))


def main():
    tam = tempfile.mkdtemp()
    dx = os.path.join(tam, "KICH-BAN-THUYET-TRINH.docx")
    tao_docx(dx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", dx, "--outdir", tam],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    shutil.copy(os.path.join(tam, "KICH-BAN-THUYET-TRINH.pdf"), PDF)
    shutil.rmtree(tam, ignore_errors=True)
    tao_md()
    print("DONE ->", PDF)
    print("DONE ->", MD)


if __name__ == "__main__":
    main()
