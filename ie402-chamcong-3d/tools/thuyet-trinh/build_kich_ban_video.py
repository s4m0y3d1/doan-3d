# -*- coding: utf-8 -*-
"""
Sinh kịch bản quay VIDEO DEMO (nộp Moodle):
  - KICH-BAN-VIDEO-DEMO.pdf
  - KICH-BAN-VIDEO-DEMO.md

Mọi con số trong kịch bản (cao độ, độ lệch, câu báo lỗi) đã được chạy thử trên
dữ liệu mẫu sạch (./chay-demo.sh --lam-moi).

    cd tools/thuyet-trinh && python3 build_kich_ban_video.py
"""
import os
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_kich_ban import CAM, GOC, NAVY, XAM, bang, phut, tieu_de

PDF = os.path.join(GOC, "KICH-BAN-VIDEO-DEMO.pdf")
MD = os.path.join(GOC, "KICH-BAN-VIDEO-DEMO.md")

CHUAN_BI = [
    ("Dữ liệu sạch. ", "Ngay trước khi quay thật: Linux chạy ./chay-demo.sh --lam-moi; Windows chạy "
     "lại CAI-DAT.bat rồi BAT-DEMO.bat. Mở http://127.0.0.1:3000/healthz phải thấy trang_thai: ok. "
     "Nếu đã tập dượt thì PHẢI làm mới lại — các con số trong kịch bản dựa trên dữ liệu mẫu ban đầu."),
    ("Trình duyệt. ", "Chrome, phóng to toàn màn hình (F11 hoặc maximize), ẩn thanh dấu trang (Ctrl+Shift+B), "
     "zoom 100 %, tắt tiện ích hiện thông báo. Mở sẵn tab http://127.0.0.1:3000 ở màn hình đăng nhập."),
    ("Dùng một trình duyệt cho cả bốn tài khoản. ", "Ở cảnh 12, nhân viên mới chấm công trên cùng trình "
     "duyệt mà Khôi vừa dùng nên hệ thống sinh cảnh báo R4 (trùng thiết bị) — đây là chủ ý để minh hoạ R4. "
     "Không mở cửa sổ ẩn danh."),
    ("Internet. ", "Bản đồ 3D và nền bản đồ tải từ máy chủ ArcGIS. Đăng nhập thử một lần rồi thoát để "
     "trình duyệt tải sẵn, tránh bản đồ trống ở đầu video."),
    ("Terminal. ", "Mở sẵn ở thư mục ie402-chamcong-3d, tăng cỡ chữ (Ctrl+Shift+= hai, ba lần) để chữ "
     "đọc được trong video."),
    ("Phần mềm quay. ", "OBS Studio: Canvas và Output 1920×1080, 30 fps, định dạng MKV (an toàn khi "
     "lỡ tắt máy) rồi File → Remux sang MP4. Nguồn: Screen Capture + Mic. Thử âm lượng mic trước."),
    ("Cách quay. ", "Quay theo từng cảnh (mỗi cảnh một tệp hoặc dùng nút Pause), đọc lời chậm, dừng 1–2 "
     "giây sau mỗi kết quả quan trọng để dễ cắt. Sau đó ghép bằng Shotcut, CapCut hoặc DaVinci Resolve."),
    ("Thao tác bản đồ. ", "Kéo chuột trái: di chuyển; kéo chuột phải: xoay / nghiêng; con lăn: phóng to "
     "thu nhỏ. Bấm vào một phần của toà nhà (khi không ở chế độ Mô phỏng) để hiện popup lớp CityGML."),
]

# (số, tên cảnh, thời lượng giây, người đọc, màn hình bắt đầu, [(thao tác, lời thoại)], kiểm tra, lưu ý)
CANH = [
    (1, "Mở đầu và khởi động hệ thống", 45, "Khôi", "Terminal ở thư mục ie402-chamcong-3d",
     [("(Tuỳ chọn) chèn slide bìa 3–5 giây khi dựng.",
       "Xin chào thầy và các bạn. Nhóm 1 xin trình bày phần demo đồ án “Hệ thống chấm công định vị 3D "
       "theo mô hình khối không gian phân tầng”."),
      ("Gõ ./chay-demo.sh --lam-moi rồi Enter. Chờ tới dòng [4/4] Sẵn sàng; trình duyệt tự mở.",
       "Toàn bộ hệ thống chạy bằng một lệnh: script dựng CSDL PostgreSQL – PostGIS, nạp lược đồ và dữ liệu "
       "mẫu, cài thư viện và bật máy chủ. Dữ liệu mẫu là toà nhà IFC One Saigon, 42 tầng, với ba công ty "
       "thuê ở ba dải tầng khác nhau.")],
     "Terminal hiện [0/4] … [4/4] Sẵn sàng; dòng “toà nhà: 1 · khối MAP: 3 · nhân viên: 4”.",
     "Trên Windows thay bằng cảnh nhấp đúp BAT-DEMO.bat."),

    (2, "Đăng nhập và mô hình 3D LOD3", 60, "Khôi", "Màn hình đăng nhập",
     [("Bấm “khoi.nm” trong bảng tài khoản thử nghiệm (tự điền mật khẩu) → bấm Đăng nhập.",
       "Em đăng nhập bằng tài khoản nhân viên Nguyễn Minh Khôi, công ty Alpha Tech, thuê tầng 5 đến 8."),
      ("Chỉ vào khung “Thông tin làm việc” bên trái.",
       "Khung bên trái là khối chấm công của nhân viên: dải cao độ 23,60 đến 42,20 mét, dung sai ngang 25 "
       "mét và dung sai đứng 5 mét."),
      ("Kéo chuột phải để xoay quanh toà nhà, lăn chuột để phóng to.",
       "Bên phải là bản đồ 3D dựng bằng ArcGIS Maps SDK. Toà nhà được dựng ở mức chi tiết 3 — có khối đế, "
       "vách kính, dầm từng tầng, khung nhôm, lối vào và mái. Ba khối trong suốt cam, xanh, tím là ba văn "
       "phòng: Alpha Tech, Beta Finance và Gamma Media."),
      ("Bấm nút LOD1 ở góc trái, chờ 2 giây, bấm lại LOD3.",
       "Nút này chuyển giữa LOD1 — khối đùn trơn, là mô hình dùng để tính toán — và LOD3 để hiển thị."),
      ("Bấm vào cửa xoay ở sảnh chính (hoặc vào vách kính, mái) → popup hiện → đóng popup (dấu ×).",
       "Mỗi thành phần gắn một lớp ngữ nghĩa theo chuẩn CityGML: cửa là Door, kính là Window, mái là "
       "RoofSurface. Toàn bộ mô hình được sinh tự động từ dữ liệu toà nhà trong CSDL.")],
     "Góc phải trên có tên Nguyễn Minh Khôi, nhãn NHAN VIEN; popup ghi “Lớp CityGML: …”.",
     "Lối vào chính ở cạnh dài nhất của khối đế, phía bờ sông Sài Gòn."),

    (3, "Tình huống 1 — đúng tầng: HỢP LỆ", 45, "Vy", "Tab Chấm công của khoi.nm",
     [("Bấm “Mô phỏng (bấm bản đồ)” → nút đổi thành “Đang chờ… bấm lên bản đồ”.",
       "Vì quay video trong phòng, em dùng chế độ mô phỏng để chọn vị trí trên bản đồ thay cho GPS."),
      ("Bấm một lần vào giữa thân toà nhà IFC One Saigon → nút đổi thành “Mô phỏng: đã chọn vị trí”.",
       "Em chọn vị trí ngay trong toà nhà."),
      ("Kéo thanh “Tầng đang làm việc” về 6. Chỉ vào tấm sàn sáng xanh lá trên bản đồ.",
       "Kéo về tầng 6: tấm sàn tầng 6 sáng màu xanh lá, nằm gọn trong dải cam của Alpha Tech."),
      ("Bấm “Chấm công VÀO”.",
       "Kết quả HỢP LỆ: khoảng cách ngang gần 0 mét, cao độ 30,58 mét, nằm trong dải 23,60 đến 42,20 mét. "
       "Ở tình huống này hàng rào 2D và hệ 3D cho cùng kết quả.")],
     "Khung xanh “HỢP LỆ — đã ghi nhận chấm công VAO”, Cao độ 30.58 m.",
     "Nếu lỡ bấm vào toà nhà khi chưa bật Mô phỏng thì popup sẽ hiện — đóng popup rồi bấm Mô phỏng."),

    (4, "Tình huống 2 — sai tầng: NGHI NGỜ (điểm nhấn)", 60, "Vy", "Tiếp cảnh 3",
     [("KHÔNG bấm lại bản đồ. Kéo thanh tầng lên 22. Dừng 2 giây, chỉ vào tấm sàn màu hổ phách.",
       "Đây là tình huống quan trọng nhất. Em giữ nguyên vị trí — cùng kinh độ, cùng vĩ độ — chỉ kéo lên "
       "tầng 22. Tấm sàn đổi màu hổ phách và nằm trong dải xanh của Beta Finance."),
      ("Bấm “Chấm công RA”. Dừng 2 giây trên khung kết quả.",
       "Kết quả NGHI NGỜ — đúng toà nhà, sai tầng. Khoảng cách ngang vẫn 0 mét, nên hàng rào 2D sẽ kết "
       "luận hợp lệ. Nhưng cao độ 104,98 mét, lệch 62,78 mét so với dải cho phép, và cảnh báo R1 mức cao "
       "cho biết lệch tương đương 13,5 tầng. Đây là trạng thái mà mô hình hai chiều không thể sinh ra.")],
     "Khung vàng “NGHI NGỜ — đúng toà nhà, sai tầng”, Lệch cao độ +62.78 m, R1 … 13.5 tầng.",
     "Đây là cảnh quan trọng nhất — nếu đọc vấp thì quay lại riêng cảnh này."),

    (5, "Tình huống 3 — ngoài toà nhà: NGOÀI VÙNG", 35, "Vy", "Tiếp cảnh 4",
     [("Bấm Mô phỏng → bấm vào một con đường cách toà nhà vài trăm mét (ví dụ đường Hàm Nghi phía trái "
       "màn hình) → kéo tầng về 6 → bấm “Chấm công VÀO”.",
       "Cuối cùng em chọn một điểm cách toà nhà khoảng hai trăm mét. Kết quả NGOÀI VÙNG vì vượt xa dung "
       "sai 25 mét. Hệ thống còn tự so với lần chấm công vừa rồi: di chuyển hai trăm mét trong vài giây là "
       "bất khả thi, nên sinh thêm cảnh báo R2.")],
     "Khung đỏ “NGOÀI VÙNG — từ chối”, khoảng cách ngang khoảng 200 m, cảnh báo R2.",
     "Khoảng cách cụ thể tuỳ chỗ bấm; R2 chỉ xuất hiện nếu bấm chấm công ngay sau cảnh 4."),

    (6, "Lịch sử chấm công và gửi giải trình", 40, "Vy", "Tiếp cảnh 5",
     [("Bấm tab “Lịch sử”.",
       "Tab lịch sử liệt kê ba lần chấm công vừa rồi, tô màu theo trạng thái."),
      ("Ở dòng NGHI_NGO (loại RA, tầng 22) bấm “Gửi giải trình” → hộp thoại hiện → gõ “Họp với khách "
       "hàng ở tầng 22” → OK → hộp thông báo → OK.",
       "Nhân viên không bị kết luận gian lận ngay — bạn ấy gửi giải trình cho bản ghi nghi ngờ."),
      ("Chỉ vào cột Giải trình hiện CHO DUYET.", "Đơn đang chờ quản lý duyệt.")],
     "Dòng NGHI_NGO có cột Giải trình = CHO DUYET.",
     "Thứ tự dòng: mới nhất ở trên — NGOAI_VUNG, NGHI_NGO, HOP_LE."),

    (7, "Quản lý: cảnh báo, duyệt đơn, báo cáo công", 75, "Vy", "Tiếp cảnh 6",
     [("Bấm Thoát → bấm “nhan.vt” → Đăng nhập → bấm tab “Quản lý”.",
       "Em chuyển sang tài khoản quản lý Võ Thành Nhân."),
      ("Chỉ lần lượt vào bảng điều khiển và bảng cảnh báo.",
       "Bảng điều khiển: một hợp lệ, một nghi ngờ, một ngoài vùng, hai cảnh báo và một đơn chờ duyệt. "
       "Cảnh báo được sắp theo mức độ trước, thời gian sau."),
      ("Ở “Đơn giải trình chờ duyệt” bấm “Duyệt”.",
       "Quản lý duyệt đơn. Bản ghi tự chuyển sang HỢP LỆ và cảnh báo R1 được đánh dấu đã xử lý — cả hai "
       "trong cùng một giao dịch CSDL."),
      ("Chỉ vào số liệu mới (Hợp lệ 2, Nghi ngờ 0) và cột “Đã xử lý” của R1 = rồi.", ""),
      ("(Tuỳ chọn) Bấm “Xuất CSV” ở Báo cáo công tháng này → mở tệp bằng Excel.",
       "Báo cáo công xuất ra CSV, mở bằng Excel đúng dấu tiếng Việt và đúng cột.")],
     "Sau khi duyệt: Hợp lệ 2 · Nghi ngờ 0 · Ngoài vùng 1 · Đơn chờ duyệt 0.",
     "Mở CSV bằng LibreOffice thì chọn dấu chấm phẩy làm phân cách ở hộp thoại nhập."),

    (8, "Quản trị: thống kê dữ liệu nền", 30, "Tín", "Tiếp cảnh 7",
     [("Bấm Thoát → bấm “tin.nh” → Đăng nhập → bấm tab “Quản trị” (mặc định ở mục Thống kê).",
       "Tài khoản quản trị Nguyễn Hữu Tín quản lý dữ liệu nền: toà nhà, văn phòng, khối MAP, nhân viên và ca."),
      ("Chỉ vào các ô số liệu; bấm dòng “Alpha Tech — Trụ sở” trong bảng → bản đồ bay tới khối.",
       "Mục thống kê cho số toà nhà, văn phòng, khối đang hiệu lực, nhân viên và tổng thể tích các khối; "
       "bảng dưới đếm lượt chấm công theo từng khối, bấm vào là bản đồ bay tới khối đó.")],
     "1 toà nhà · 3 văn phòng · 3 khối · 4 nhân viên · 3 công ty · 2 ca.",
     "Bản đồ quản trị dùng nền ảnh vệ tinh."),

    (9, "Quản trị: thêm toà nhà bằng cách vẽ trên bản đồ", 75, "Tín", "Tab Quản trị của tin.nh",
     [("Kéo / lăn bản đồ tới một khu đất trống gần IFC One Saigon (phía đường Hàm Nghi).",
       "Em thêm một toà nhà mới."),
      ("Bấm mục “Toà nhà” → “+ Thêm mới”. Nhập: Tên “Saigon Centre Demo”, Địa chỉ “Đường Hàm Nghi, Phường "
       "Sài Gòn”, Cao độ nền 5, Số tầng 18, Chiều cao tầng 3.8, Số tầng khối đế 3.",
       "Nhập tên, cao độ nền, số tầng, chiều cao tầng và số tầng khối đế."),
      ("Bấm “Vẽ đa giác trên bản đồ” → bấm 4 điểm tạo một hình tứ giác khoảng 40–50 m → bấm đúp ở điểm "
       "cuối để kết thúc.",
       "Mặt bằng toà nhà được vẽ trực tiếp trên bản đồ 3D. Vẽ xong là mô hình LOD3 hiện ngay để xem trước: "
       "khối đế ba tầng, thân tháp, dầm tầng, lối vào và mái — sinh tự động từ thông số vừa nhập."),
      ("(Tuỳ chọn) Bấm “Sửa đỉnh”, kéo một đỉnh, bấm ra ngoài.", "Có thể chỉnh lại từng đỉnh."),
      ("Bấm “Lưu”.", "Lưu lại — danh sách giờ có hai toà nhà.")],
     "Bản xem trước LOD3 có nhãn “Saigon Centre Demo (xem trước)”; sau khi lưu: “Toà nhà (2)”.",
     "Đa giác không được cắt chéo chính nó. Nhấn Esc để huỷ nét vẽ dở."),

    (10, "Quản trị: thêm văn phòng và khối MAP — ràng buộc CSDL", 60, "Tín", "Tiếp cảnh 9",
     [("Bấm mục “Văn phòng & khối MAP” → “+ Thêm mới”. Nhập Tên “Alpha Tech — Chi nhánh 2”, Công ty "
       "thuê “Công ty TNHH Alpha Tech”, Toà nhà “Saigon Centre Demo (18 tầng)”, Từ tầng 9, Đến tầng 20.",
       "Em cho Alpha Tech thuê thêm một văn phòng ở toà nhà mới. Thử nhập đến tầng 20, trong khi toà "
       "nhà chỉ có 18 tầng."),
      ("Chỉ vào khối xem trước màu đỏ vượt quá mái. Bấm “Lưu”.",
       "Khối xem trước báo đỏ vì vượt mái. Bấm lưu thì CSDL từ chối bằng trigger, kèm thông báo tiếng "
       "Việt — dù sửa bằng giao diện hay bằng SQL cũng không lọt."),
      ("Sửa “Đến tầng” thành 12 → khối chuyển màu vàng, gợi ý “z = 35.40 → 50.60 m” → bấm “Lưu”.",
       "Sửa lại tầng 9 đến 12: cao độ khối 35,40 đến 50,60 mét được tính ngay, và giá trị lưu thật do "
       "trigger trong CSDL suy ra từ cao độ nền và chiều cao tầng.")],
     "Dòng đỏ “Toà nhà chỉ có 18 tầng, không thể gán văn phòng tới tầng 20”; sau khi lưu bảng có dòng 9–12 · 35.4–50.6.",
     ""),

    (11, "Quản trị: nhân viên, ca làm việc, phân ca", 60, "Tín", "Tiếp cảnh 10",
     [("Bấm mục “Nhân viên” → “+ Thêm mới”. Nhập Họ tên “Lê Văn Demo”, Email demo.lv@alphatech.vn, Công ty "
       "Alpha Tech, Phòng ban “Phòng Kỹ thuật”, Tên đăng nhập demo.lv, Mật khẩu 123456, Vai trò Nhân viên → Lưu.",
       "Thêm một nhân viên mới cho Alpha Tech, tạo luôn tài khoản đăng nhập. Phòng ban được lọc theo công ty."),
      ("Bấm mục “Ca làm việc” → “+ Thêm mới” → Tên “Ca tối”, Giờ vào 18:00, Giờ ra 22:00, Trễ 10 → Lưu.",
       "Thêm một ca làm việc."),
      ("Ở khung “Phân ca cho nhân viên”: chọn Lê Văn Demo, ca Hành chính, giữ khoảng ngày mặc định → bấm "
       "“Gán ca”.",
       "Và phân ca hành chính cho nhân viên mới trong bảy ngày tới, mặc định bỏ qua cuối tuần.")],
     "Gợi ý dưới bản đồ “Đã gán ca cho … ngày”; lịch 7 ngày có tên Lê Văn Demo.",
     "Ô giờ hiển thị theo ngôn ngữ máy (có thể là 06:00 PM) — vẫn đúng. Số ngày được gán tuỳ ngày quay."),

    (12, "Nhân viên mới chấm công ở toà nhà vừa dựng", 45, "Nhân", "Tiếp cảnh 11",
     [("Bấm Thoát → gõ tay demo.lv / 123456 → Đăng nhập.",
       "Đăng nhập bằng nhân viên vừa tạo."),
      ("Chỉ vào dòng Văn phòng “Alpha Tech — Trụ sở (+1 khối khác)”.",
       "Công ty của bạn ấy có hai khối: trụ sở ở IFC One Saigon và chi nhánh ở toà nhà mới."),
      ("Bấm Mô phỏng → bấm vào thân toà “Saigon Centre Demo” → kéo tầng về 10 → “Chấm công VÀO”.",
       "Chấm công ở tầng 10 toà nhà mới: hệ thống tự chọn đúng khối gần nhất theo vị trí và tầng, kết quả "
       "HỢP LỆ, cao độ 41,10 mét trong dải 35,40 đến 50,60."),
      ("Chỉ vào cảnh báo R4.",
       "Có thêm cảnh báo R4 — trùng thiết bị — vì cùng trình duyệt này vừa được bạn Khôi dùng để chấm công "
       "chưa tới mười phút. Đây là quy tắc chống chấm công hộ.")],
     "Khung xanh HỢP LỆ, Cao độ 41.1 m, dải 35.40–50.60, cảnh báo R4.",
     "Toà nhà mới nằm cạnh IFC One Saigon; xoay / lăn bản đồ nếu chưa thấy."),

    (13, "Quản trị: sửa và xoá", 60, "Nhân", "Tiếp cảnh 12",
     [("Thoát → tin.nh → Quản trị → mục “Toà nhà” → bấm “Sửa” ở Saigon Centre Demo → Chiều cao tầng 4 "
       "→ Lưu → bấm mục “Văn phòng & khối MAP”.",
       "Sửa chiều cao tầng toà nhà từ 3,8 lên 4 mét. Khối của chi nhánh 2 tự tính lại thành 37 đến 53 mét — "
       "trigger trong CSDL cập nhật mọi khối bên trong toà nhà."),
      ("Mục “Toà nhà” → bấm “Xoá” ở Saigon Centre Demo → OK ở hộp xác nhận → đọc thông báo → OK.",
       "Thử xoá toà nhà: không được, vì toà nhà còn văn phòng — khoá ngoại bảo vệ dữ liệu."),
      ("Mục “Nhân viên” → “Xoá” ở Lê Văn Demo → OK → đọc thông báo và gợi ý → OK.",
       "Nhân viên đã có bản ghi chấm công cũng không xoá cứng được, để giữ lịch sử công. Hệ thống gợi ý "
       "chuyển sang nghỉ việc."),
      ("Bấm “Sửa” ở Lê Văn Demo → bỏ tick “Đang làm việc” → Lưu.",
       "Bỏ chọn đang làm việc: tài khoản bị khoá đăng nhập nhưng lịch sử vẫn còn."),
      ("Mục “Ca làm việc” → “Xoá” ở Ca tối → OK.", "Ca chưa được dùng thì xoá bình thường.")],
     "Khối chi nhánh 2: 37.0–53.0; dòng Lê Văn Demo có nhãn “nghỉ việc”; Ca tối biến mất.",
     ""),

    (14, "Kiểm thử tự động", 30, "Nhân", "Terminal",
     [("Gõ cd server && npm test → Enter. Chờ dòng cuối.",
       "Cuối cùng là bộ kiểm thử tự động gồm 34 ca, gọi API thật trên CSDL thật: ba tình huống của mô hình, "
       "phân quyền, thêm sửa xoá, các ràng buộc và việc trải khối theo mô hình SSM. Kết quả 34 trên 34 đạt.")],
     "Dòng cuối “Kết quả: 34 đạt, 0 trượt”.",
     "Bộ kiểm thử tự dọn dữ liệu nó tạo ra; không ảnh hưởng dữ liệu vừa demo."),

    (15, "Kết thúc", 20, "Vy", "Trình duyệt, tab Chấm công, góc nhìn toàn cảnh hai toà nhà",
     [("Xoay chậm bản đồ quanh hai toà nhà LOD3 (có thể chèn slide cảm ơn khi dựng).",
       "Tóm lại, khi chuyển phép kiểm tra từ “điểm thuộc đa giác” sang “điểm thuộc khối”, hệ thống phát "
       "hiện được trạng thái đúng toà nhà nhưng sai tầng — điều hàng rào hai chiều không làm được. Cảm ơn "
       "thầy và các bạn đã theo dõi.")],
     "", "Sau khi quay xong: chạy ./chay-demo.sh --lam-moi để đưa dữ liệu về ban đầu."),
]

SAU_KHI_QUAY = [
    "Cắt bỏ khoảng chờ tải bản đồ và các đoạn đọc vấp; giữ nguyên các khung kết quả ít nhất 2 giây.",
    "Thêm chú thích chữ (text overlay) ở 3 kết quả chính: HỢP LỆ, NGHI NGỜ (lệch 13,5 tầng), NGOÀI VÙNG.",
    "Xuất MP4, H.264, 1920×1080, 30 fps; đặt tên IE402_Nhom1_DemoChamCong3D.mp4.",
    "Xem lại toàn bộ video một lần với âm thanh trước khi nộp Moodle.",
    "Chạy ./chay-demo.sh --lam-moi để dữ liệu trở về trạng thái mẫu.",
]

SU_CO = [
    ["Bản đồ trống / chỉ có nền xám", "Kiểm tra Internet; tải lại trang (F5) một lần, đăng nhập lại."],
    ["Bấm bản đồ khi Mô phỏng mà không nhận", "Nút phải đang hiện “Đang chờ… bấm lên bản đồ”; nếu popup che, đóng popup rồi bấm lại."],
    ["Kết quả khác kịch bản (ví dụ có thêm R2, R4 ở cảnh 3)", "Dữ liệu còn bản ghi của lần tập — chạy ./chay-demo.sh --lam-moi rồi quay lại từ cảnh 2."],
    ["Đăng nhập báo sai mật khẩu ở cảnh 12", "Kiểm tra đã lưu nhân viên ở cảnh 11 và gõ đúng demo.lv / 123456."],
    ["Vẽ đa giác xong không thấy bản xem trước", "Kiểm tra đã nhập Số tầng và Chiều cao tầng; bấm “Vẽ đa giác trên bản đồ” và vẽ lại."],
    ["Cổng 3000 bận khi chạy script", "./tat-demo.sh rồi chạy lại ./chay-demo.sh --lam-moi."],
]


def doan(doc, text, bold=False, mau=None, co=None, nghieng=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = nghieng
    if mau:
        r.font.color.rgb = mau
    if co:
        r.font.size = Pt(co)
    return p


def tao_docx(duong):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.8)
        sec.left_margin = sec.right_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(11)
    for ten in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[ten].font.name = "Cambria"

    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Kịch bản quay video demo · IE402 · Nhóm 1 · trang ")
    r.font.size = Pt(9)
    r.font.color.rgb = XAM
    for kieu, chu in (("begin", None), (None, "PAGE"), ("end", None)):
        rr = p.add_run()
        rr.font.size = Pt(9)
        if kieu:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kieu)
            rr._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = chu
            rr._r.append(it)

    t = doan(doc, "KỊCH BẢN QUAY VIDEO DEMO", bold=True, mau=NAVY, co=24)
    t.runs[0].font.name = "Cambria"
    doan(doc, "Hệ thống chấm công định vị 3D theo mô hình khối không gian phân tầng", mau=CAM, co=14)
    tong = sum(c[2] for c in CANH)
    p = doc.add_paragraph()
    p.add_run("Thời lượng dự kiến: ").bold = True
    p.add_run("khoảng %d phút (%d cảnh). " % (round(tong / 60), len(CANH)))
    p.add_run("Mục đích: ").bold = True
    p.add_run("video demo nộp Moodle kèm báo cáo. Mỗi cảnh gồm thao tác (làm gì trên màn hình), lời "
              "thoại (đọc gì) và điều phải thấy trên màn hình để biết đã quay đúng. Các con số trong lời "
              "thoại đã được chạy thử trên dữ liệu mẫu sạch.")
    p = doc.add_paragraph()
    p.add_run("Người đọc lời: ").bold = True
    p.add_run("gợi ý chia theo phân công (Khôi — bản đồ 3D, Vy — giao diện, Tín — dữ liệu nền, Nhân — "
              "máy chủ và kiểm thử). Nếu chỉ một người đọc thì bỏ qua cột này.")

    tieu_de(doc, "1. Chuẩn bị trước khi quay", 1)
    for dam, thuong in CHUAN_BI:
        q = doc.add_paragraph(style="List Bullet")
        q.add_run(dam).bold = True
        q.add_run(thuong)
    doan(doc, "Tài khoản dùng trong video (mật khẩu đều là 123456): khoi.nm — nhân viên Alpha Tech; "
              "nhan.vt — quản lý; tin.nh — quản trị; demo.lv — tạo mới ở cảnh 11.", nghieng=True)

    tieu_de(doc, "2. Tổng quan các cảnh", 1)
    da = 0
    dong = []
    for so, ten, giay, nguoi, *_ in CANH:
        dong.append([str(so), ten, "%s → %s" % (phut(da), phut(da + giay)), nguoi])
        da += giay
    bang(doc, ["#", "Cảnh", "Mốc thời gian", "Người đọc"], dong, [1.0, 10.2, 3.2, 2.2])

    tieu_de(doc, "3. Chi tiết từng cảnh", 1)
    da = 0
    for so, ten, giay, nguoi, bat_dau, buoc, kiem_tra, luu_y in CANH:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.keep_with_next = True
        r = h.add_run("Cảnh %d · %s" % (so, ten))
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = NAVY
        r.font.name = "Cambria"
        m = doc.add_paragraph()
        m.paragraph_format.keep_with_next = True
        r = m.add_run("%s → %s  ·  ~%s  ·  người đọc: %s  ·  bắt đầu từ: %s"
                      % (phut(da), phut(da + giay), phut(giay), nguoi, bat_dau))
        r.font.size = Pt(9.5)
        r.font.color.rgb = XAM
        da += giay
        bang(doc, ["#", "Thao tác trên màn hình", "Lời thoại"],
             [[str(k), tt, lt or "—"] for k, (tt, lt) in enumerate(buoc, start=1)], [0.8, 7.4, 8.8])
        if kiem_tra:
            q = doc.add_paragraph()
            q.add_run("✔ Trên màn hình phải thấy: ").bold = True
            q.add_run(kiem_tra)
        if luu_y:
            q = doc.add_paragraph()
            rr = q.add_run("Lưu ý: ")
            rr.bold = True
            rr.font.color.rgb = CAM
            q.add_run(luu_y)

    tieu_de(doc, "4. Xử lý sự cố khi quay", 1)
    bang(doc, ["Tình huống", "Cách xử lý"], SU_CO, [5.6, 11.4])

    tieu_de(doc, "5. Sau khi quay", 1)
    for c in SAU_KHI_QUAY:
        doc.add_paragraph(c, style="List Bullet")
    doc.save(duong)


def tao_md():
    d = ["# Kịch bản quay video demo — Đồ án IE402", "",
         "> Sinh tự động từ `tools/thuyet-trinh/build_kich_ban_video.py` (bản in: `KICH-BAN-VIDEO-DEMO.pdf`).", "",
         "## Chuẩn bị", ""]
    d += ["- **%s**%s" % (a.strip(), " " + b) for a, b in CHUAN_BI] + [""]
    da = 0
    d += ["## Các cảnh", ""]
    for so, ten, giay, nguoi, bat_dau, buoc, kiem_tra, luu_y in CANH:
        d += ["### Cảnh %d · %s (%s → %s, %s)" % (so, ten, phut(da), phut(da + giay), nguoi), "",
              "_Bắt đầu từ: %s_" % bat_dau, "", "| Bước | Thao tác | Lời thoại |", "|---|---|---|"]
        d += ["| %d | %s | %s |" % (k, tt, lt or "—") for k, (tt, lt) in enumerate(buoc, start=1)]
        d += [""]
        if kiem_tra:
            d += ["**Phải thấy:** " + kiem_tra, ""]
        if luu_y:
            d += ["**Lưu ý:** " + luu_y, ""]
        da += giay
    d += ["## Xử lý sự cố", "", "| Tình huống | Cách xử lý |", "|---|---|"]
    d += ["| %s | %s |" % tuple(r) for r in SU_CO] + ["", "## Sau khi quay", ""]
    d += ["- " + c for c in SAU_KHI_QUAY] + [""]
    open(MD, "w", encoding="utf-8").write("\n".join(d))


def main():
    tam = tempfile.mkdtemp()
    dx = os.path.join(tam, "KICH-BAN-VIDEO-DEMO.docx")
    tao_docx(dx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", dx, "--outdir", tam],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    shutil.copy(os.path.join(tam, "KICH-BAN-VIDEO-DEMO.pdf"), PDF)
    shutil.rmtree(tam, ignore_errors=True)
    tao_md()
    print("DONE ->", PDF)
    print("DONE ->", MD)


if __name__ == "__main__":
    main()
