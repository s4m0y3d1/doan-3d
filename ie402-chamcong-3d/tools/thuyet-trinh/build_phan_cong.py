# -*- coding: utf-8 -*-
"""
Sinh BẢNG PHÂN CÔNG VÀ ĐÁNH GIÁ THÀNH VIÊN (nộp Moodle):
  - PHAN-CONG-DANH-GIA-NHOM.pdf

Danh sách thành viên và công việc lấy từ noi-dung.json (dùng chung với slide).

    cd tools/thuyet-trinh && python3 build_phan_cong.py
"""
import os
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_kich_ban import CAM, GOC, ND, NAVY, bang, tieu_de

PDF = os.path.join(GOC, "PHAN-CONG-DANH-GIA-NHOM.pdf")
DANH_GIA = ("Đạt", "10")


def giua(doc, text, bold=False, co=None, mau=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    if co:
        r.font.size = Pt(co)
    if mau:
        r.font.color.rgb = mau
    p.paragraph_format.space_after = Pt(2)
    return p


def tao_docx(duong):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(1.6)
        sec.left_margin = sec.right_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(11)
    for ten in ("Heading 1", "Heading 2"):
        doc.styles[ten].font.name = "Cambria"

    giua(doc, "ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH", co=12)
    giua(doc, "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN", bold=True, co=12)
    t = giua(doc, "BẢNG PHÂN CÔNG VÀ ĐÁNH GIÁ THÀNH VIÊN", bold=True, co=18, mau=NAVY)
    t.runs[0].font.name = "Cambria"
    giua(doc, ND["tieu_de"], co=12.5, mau=CAM).paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run("Môn học: ").bold = True
    p.add_run(ND["mon"])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run("Giảng viên hướng dẫn: ").bold = True
    p.add_run(ND["gvhd"])
    p = doc.add_paragraph()
    p.add_run("Nhóm: ").bold = True
    p.add_run(ND["nhom"] + " — nhóm trưởng: " + ND["thanh_vien"][0]["ten"])

    tieu_de(doc, "1. Bảng phân công công việc", 1).paragraph_format.space_before = Pt(8)
    bang(doc, ["STT", "MSSV", "Họ và tên", "Vai trò", "Công việc phụ trách"],
         [[str(i), tv["mssv"], tv["ten"], tv["vai"], tv["viec"]]
          for i, tv in enumerate(ND["thanh_vien"], start=1)],
         [1.1, 2.2, 3.3, 3.2, 7.2])
    p = doc.add_paragraph()
    p.add_run("Nguyên tắc chia việc: ").bold = True
    p.add_run("chia theo lớp kiến trúc của hệ thống; mỗi phần có một người chính và một người đọc chéo; "
              "nhóm chốt lược đồ CSDL và danh sách API sớm để bốn thành viên làm song song.")

    tieu_de(doc, "2. Đánh giá thành viên", 1).paragraph_format.space_before = Pt(8)
    bang(doc, ["STT", "MSSV", "Họ và tên", "Mức độ hoàn thành", "Điểm (thang 10)"],
         [[str(i), tv["mssv"], tv["ten"], DANH_GIA[0], DANH_GIA[1]]
          for i, tv in enumerate(ND["thanh_vien"], start=1)],
         [1.1, 2.4, 4.6, 4.4, 4.5])
    doc.add_paragraph("Nhóm thống nhất đánh giá tất cả thành viên: Đạt, 10/10 điểm.")

    doc.save(duong)


def main():
    tam = tempfile.mkdtemp()
    dx = os.path.join(tam, "PHAN-CONG-DANH-GIA-NHOM.docx")
    tao_docx(dx)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", dx, "--outdir", tam],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    shutil.copy(os.path.join(tam, "PHAN-CONG-DANH-GIA-NHOM.pdf"), PDF)
    shutil.rmtree(tam, ignore_errors=True)
    print("DONE ->", PDF)


if __name__ == "__main__":
    main()
